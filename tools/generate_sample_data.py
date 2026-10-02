"""DEV TOOL: generate the four monthly sales workbooks used to test the app.

This is NOT part of the application. It simply produces realistic input files
that the app will later detect and consolidate into Master_Sales_Report.xlsx.

Columns A-J hold typed values; columns K-N are real Excel formulas.

Run from the project root:
    python tools\\generate_sample_data.py
"""
from __future__ import annotations

import argparse
import calendar
import random
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "data" / "input"

YEAR = 2026
MONTHS = [(1, "January"), (2, "February"), (3, "March"), (4, "April")]
SHEET_NAME = "Sales Data"

COLUMNS = [
    "Date", "Invoice No.", "Customer", "Product", "Category", "Quantity",
    "Unit Price", "Discount", "Tax Rate", "Salesperson",
    "Gross Amount", "Discount Amount", "Tax Amount", "Net Sales",
]
# F=Quantity G=Unit Price H=Discount I=Tax Rate K=Gross L=DiscAmt M=Tax N=Net
FORMULAS = [
    "=F{r}*G{r}",
    "=K{r}*H{r}",
    "=(K{r}-L{r})*I{r}",
    "=K{r}-L{r}+M{r}",
]

CUSTOMERS = [
    "Acacia Retail Ltd", "Baobab Traders", "Cedar Hardware", "Delta Office Supplies",
    "Equator Logistics", "Fig Tree Cafe", "Green Valley Schools", "Highland Clinic",
    "Ivory Interiors", "Jacaranda Hotels", "Kilimani Tech Hub", "Lakeview Pharmacy",
    "Mara Safari Outfitters", "Nile Stationers", "Orchid Salon", "Pioneer Motors",
    "Quarry Construction", "Rift Valley Farms", "Savanna Bookshop", "Tembo Supermarket",
    "Umoja Sacco", "Victoria Textiles", "Wema Foundation", "Zebra Print Studio",
]
# (product, category, min price, max price)
PRODUCTS = [
    ("Laptop Pro 15", "Electronics", 780.00, 1250.00),
    ("Wireless Mouse", "Electronics", 12.00, 35.00),
    ("Mechanical Keyboard", "Electronics", 45.00, 130.00),
    ("27-inch Monitor", "Electronics", 180.00, 420.00),
    ("Ergonomic Chair", "Furniture", 140.00, 380.00),
    ("Standing Desk", "Furniture", 260.00, 640.00),
    ("Filing Cabinet", "Furniture", 90.00, 210.00),
    ("A4 Paper Ream", "Stationery", 3.50, 6.50),
    ("Ballpoint Pens (Box)", "Stationery", 4.00, 9.00),
    ("Desk Organizer", "Stationery", 8.00, 22.00),
    ("Whiteboard Markers", "Stationery", 5.00, 12.00),
    ("Laser Printer", "Electronics", 210.00, 520.00),
    ("Printer Toner", "Consumables", 38.00, 95.00),
    ("Coffee Beans 1kg", "Consumables", 11.00, 24.00),
    ("Bottled Water (Case)", "Consumables", 6.00, 12.00),
]
SALESPERSONS = [
    "Amina Otieno", "Brian Mwangi", "Carol Njeri", "David Kamau",
    "Esther Wanjiru", "Felix Odhiambo", "Grace Achieng", "Hassan Ali",
]
DISCOUNT_RATES = [0.00, 0.05, 0.10, 0.15]
DISCOUNT_WEIGHTS = [50, 25, 15, 10]
TAX_BY_CATEGORY = {"Electronics": 0.16, "Furniture": 0.16, "Stationery": 0.08, "Consumables": 0.00}

# --- Styling ------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FORMATS = {
    "Date": "yyyy-mm-dd", "Quantity": "#,##0", "Unit Price": "$#,##0.00",
    "Discount": "0%", "Tax Rate": "0%", "Gross Amount": "$#,##0.00",
    "Discount Amount": "$#,##0.00", "Tax Amount": "$#,##0.00", "Net Sales": "$#,##0.00",
}
WIDTHS = [12, 18, 26, 24, 14, 10, 12, 10, 10, 18, 15, 16, 13, 15]


def build_base_rows(month: int, rows: int, rng: random.Random) -> list[list]:
    """Typed (non-formula) values for one month, sorted by date."""
    days = calendar.monthrange(YEAR, month)[1]
    dates = sorted(date(YEAR, month, rng.randint(1, days)) for _ in range(rows))
    out = []
    for seq, d in enumerate(dates, start=1):
        product, category, low, high = rng.choice(PRODUCTS)
        out.append([
            datetime(d.year, d.month, d.day),
            f"INV-{YEAR}{month:02d}-{seq:04d}",
            rng.choice(CUSTOMERS),
            product,
            category,
            rng.randint(1, 25),
            round(rng.uniform(low, high), 2),
            rng.choices(DISCOUNT_RATES, weights=DISCOUNT_WEIGHTS)[0],
            TAX_BY_CATEGORY[category],
            rng.choice(SALESPERSONS),
        ])
    return out


def create_workbook(month: int, rows: int, rng: random.Random, dest: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    ws.append(COLUMNS)
    for i, name in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS[i - 1]
    ws.row_dimensions[1].height = 24

    for offset, base in enumerate(build_base_rows(month, rows, rng)):
        r = offset + 2
        ws.append(base + [f.format(r=r) for f in FORMULAS])
        for i, name in enumerate(COLUMNS, start=1):
            c = ws.cell(row=r, column=i)
            c.border = BORDER
            if name in FORMATS:
                c.number_format = FORMATS[name]

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{rows + 1}"
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)


def main() -> None:
    p = argparse.ArgumentParser(description="Generate sample monthly sales workbooks.")
    p.add_argument("--rows", type=int, default=100, help="data rows per file (min 50)")
    p.add_argument("--seed", type=int, default=42)
    p.add_argument("--out", type=Path, default=DEFAULT_OUT, help="output folder")
    a = p.parse_args()
    if a.rows < 50:
        raise SystemExit("--rows must be at least 50")

    rng = random.Random(a.seed)
    for month, name in MONTHS:
        dest = a.out / f"Sales_{name}.xlsx"
        create_workbook(month, a.rows, rng, dest)
        print(f"Created {dest}  ({a.rows} rows x {len(COLUMNS)} columns)")


if __name__ == "__main__":
    main()