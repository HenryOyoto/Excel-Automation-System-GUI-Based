"""DEV TOOL: independently verify Master_Sales_Report.xlsx against the raw inputs.

It recomputes every number straight from the input workbooks with openpyxl only,
so a bug in the app's engine cannot hide itself.

Run from the project root:
    python tools\\verify_master.py
"""
from __future__ import annotations

import math
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "data" / "input"
MASTER = ROOT / "data" / "output" / "Master_Sales_Report.xlsx"

SHEETS = ["Dashboard", "Consolidated Data", "Product Analysis",
          "Salesperson Analysis", "Monthly Summary"]
MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]
CONSOL_HEADERS = ["Date", "Invoice No.", "Customer", "Product", "Category", "Quantity",
                  "Unit Price", "Discount", "Tax Rate", "Salesperson", "Gross Amount",
                  "Discount Amount", "Tax Amount", "Net Sales", "Source File", "Month"]

results: list[bool] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    results.append(ok)
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f"   -> {detail}" if detail and not ok else ""))


def close(a, b) -> bool:
    return a is not None and b is not None and math.isclose(float(a), float(b), rel_tol=1e-9, abs_tol=0.01)


def load_expected():
    files = sorted(p for p in INPUT.glob("*.xlsx")
                   if not p.name.startswith("~$") and p.name != MASTER.name)
    recs = []
    for p in files:
        ws = load_workbook(p).worksheets[0]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in row):
                continue
            d, inv, cust, prod, cat, qty, price, disc, tax, sp = row[:10]
            gross = qty * price
            damt = gross * disc
            tamt = (gross - damt) * tax
            recs.append(dict(date=d, invoice=inv, customer=cust, product=prod, category=cat,
                             qty=qty, price=price, sp=sp, gross=gross, damt=damt, tamt=tamt,
                             net=gross - damt + tamt, source=p.name))
    return files, recs


def dash_value(ws, label):
    for row in ws.iter_rows():
        for c in row:
            if c.value == label:
                return ws.cell(row=c.row + 1, column=c.column).value
    return None


def main() -> int:
    files, exp = load_expected()
    print(f"Inputs: {len(files)} files, {len(exp)} rows\nMaster: {MASTER}\n")
    if not MASTER.exists() or not exp:
        print("Master report or input files missing.")
        return 1
    try:
        wb = load_workbook(MASTER, data_only=True)
    except Exception as exc:
        print(f"[FAIL] master does not open: {exc}")
        return 1
    check("Master workbook opens", True)
    check("Five sheets in the right order", wb.sheetnames == SHEETS, str(wb.sheetnames))
    if wb.sheetnames != SHEETS:
        return 1

    # ---- Consolidated Data ----
    ws = wb["Consolidated Data"]
    headers = [c.value for c in ws[1]]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[1] is not None]
    check("Consolidated headers correct", headers == CONSOL_HEADERS, str(headers))
    check("Consolidated row count == total input rows", len(rows) == len(exp), f"{len(rows)} vs {len(exp)}")
    check("Invoice numbers match inputs exactly",
          Counter(r[1] for r in rows) == Counter(e["invoice"] for e in exp))
    check("Rows per source file match",
          Counter(r[14] for r in rows) == Counter(e["source"] for e in exp))
    by_inv = {e["invoice"]: e for e in exp}
    bad = []
    for r in rows:
        e = by_inv.get(r[1])
        ok = (e is not None and isinstance(r[0], datetime) and r[0].date() == e["date"].date()
              and close(r[10], e["gross"]) and close(r[11], e["damt"])
              and close(r[12], e["tamt"]) and close(r[13], e["net"])
              and r[14] == e["source"] and str(r[15]).split()[0] == MONTHS[r[0].month - 1])
        if not ok:
            bad.append(r[1])
    check("Every row: date, gross, discount, tax, net, source, month correct", not bad, str(bad[:5]))

    # ---- Product Analysis ----
    ws = wb["Product Analysis"]
    exp_p = defaultdict(lambda: [0, 0, 0.0, 0.0])
    for e in exp:
        g = exp_p[(e["product"], e["category"])]
        g[0] += 1; g[1] += e["qty"]; g[2] += e["net"]; g[3] += e["price"]
    got = {(r[0], r[1]): r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None}
    check("Product Analysis has every product", set(got) == set(exp_p))
    check("Product Analysis values correct",
          all(k in got and got[k][2] == g[0] and got[k][3] == g[1] and close(got[k][4], g[2])
              and close(got[k][5], g[3] / g[0]) for k, g in exp_p.items()))

    # ---- Salesperson Analysis ----
    ws = wb["Salesperson Analysis"]
    exp_s = defaultdict(lambda: [0, 0, 0.0])
    for e in exp:
        g = exp_s[e["sp"]]
        g[0] += 1; g[1] += e["qty"]; g[2] += e["net"]
    got = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None}
    check("Salesperson Analysis has every salesperson", set(got) == set(exp_s))
    check("Salesperson Analysis values correct",
          all(k in got and got[k][1] == g[0] and got[k][2] == g[1] and close(got[k][3], g[2])
              and close(got[k][4], g[2] / g[0]) for k, g in exp_s.items()))

    # ---- Monthly Summary ----
    ws = wb["Monthly Summary"]
    exp_m = defaultdict(lambda: [0, 0, 0.0, 0.0, 0.0])
    for e in exp:
        g = exp_m[e["date"].month]
        g[0] += 1; g[1] += e["qty"]; g[2] += e["net"]; g[3] += e["damt"]; g[4] += e["tamt"]
    got = {MONTHS.index(str(r[0]).split()[0]) + 1: r
           for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None}
    names = ", ".join(MONTHS[m - 1] for m in sorted(exp_m))
    check(f"Monthly Summary contains all months ({names})", set(got) == set(exp_m))
    check("Monthly Summary values correct",
          all(m in got and got[m][1] == g[0] and got[m][2] == g[1] and close(got[m][3], g[2])
              and close(got[m][4], g[3]) and close(got[m][5], g[4]) and close(got[m][6], g[2] / g[0])
              for m, g in exp_m.items()))

    # ---- Dashboard ----
    ws = wb["Dashboard"]
    total_net = sum(e["net"] for e in exp)
    expected_dash = {
        "Total Revenue": total_net,
        "Total Quantity": sum(e["qty"] for e in exp),
        "Total Tax": sum(e["tamt"] for e in exp),
        "Total Discount": sum(e["damt"] for e in exp),
        "Average Order Value": total_net / len(exp),
        "Number of Transactions": len(exp),
        "Number of Customers": len({e["customer"] for e in exp}),
        "Number of Products": len({e["product"] for e in exp}),
        "Number of Salespersons": len({e["sp"] for e in exp}),
    }
    for label, want in expected_dash.items():
        got_v = dash_value(ws, label)
        check(f"Dashboard '{label}' = {want:,.2f}", close(got_v, want), f"found {got_v}")

    # ---- Structure / formatting ----
    with zipfile.ZipFile(MASTER) as z:
        n_charts = len([n for n in z.namelist() if n.startswith("xl/charts/chart")])
    check("Dashboard has 3 charts", n_charts == 3, f"found {n_charts}")
    wbf = load_workbook(MASTER)
    for s in SHEETS[1:]:
        w = wbf[s]
        check(f"'{s}': frozen header and Excel table",
              w.freeze_panes == "A2" and len(w.tables) == 1)
    for s in SHEETS[2:]:
        check(f"'{s}': conditional formatting present",
              len(list(wbf[s].conditional_formatting)) > 0)

    failed = results.count(False)
    print(f"\n{results.count(True)} passed, {failed} failed")
    print("RESULT: FAILED" if failed else "RESULT: MASTER WORKBOOK VERIFIED")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())