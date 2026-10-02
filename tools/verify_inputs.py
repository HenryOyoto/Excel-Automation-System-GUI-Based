"""DEV TOOL: verify the generated input workbooks before the app touches them.

Run from the project root:
    python tools\\verify_inputs.py
"""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / "data" / "input"
FILES = [f"Sales_{m}.xlsx" for m in ("January", "February", "March", "April")]
SHEET = "Sales Data"

COLUMNS = [
    "Date", "Invoice No.", "Customer", "Product", "Category", "Quantity",
    "Unit Price", "Discount", "Tax Rate", "Salesperson",
    "Gross Amount", "Discount Amount", "Tax Amount", "Net Sales",
]
FORMULAS = ["=F{r}*G{r}", "=K{r}*H{r}", "=(K{r}-L{r})*I{r}", "=K{r}-L{r}+M{r}"]
MIN_ROWS, MIN_COLS = 50, 10


def verify(path: Path) -> tuple[list[str], int]:
    if not path.exists():
        return [f"file not found: {path}"], 0
    try:
        wb = load_workbook(path)
    except Exception as exc:
        return [f"cannot open workbook: {exc}"], 0
    if SHEET not in wb.sheetnames:
        return [f"missing sheet '{SHEET}'"], 0

    ws = wb[SHEET]
    problems: list[str] = []
    headers = [c.value for c in ws[1]]
    if len(headers) < MIN_COLS:
        problems.append(f"only {len(headers)} columns (min {MIN_COLS})")
    if headers != COLUMNS:
        problems.append(f"header mismatch: {headers}")

    rows = ws.max_row - 1
    if rows < MIN_ROWS:
        problems.append(f"only {rows} data rows (min {MIN_ROWS})")

    invoices = set()
    for r in range(2, ws.max_row + 1):
        v = [c.value for c in ws[r]]
        for i in range(10):
            if v[i] in (None, ""):
                problems.append(f"row {r}: empty '{COLUMNS[i]}'")
        if not isinstance(v[0], datetime):
            problems.append(f"row {r}: invalid date {v[0]!r}")
        for i in (5, 6, 7, 8):
            if not isinstance(v[i], (int, float)) or isinstance(v[i], bool):
                problems.append(f"row {r}: '{COLUMNS[i]}' not numeric ({v[i]!r})")
        for i, tmpl in enumerate(FORMULAS, start=10):
            expected = tmpl.format(r=r)
            if v[i] != expected:
                problems.append(f"row {r}: '{COLUMNS[i]}' is {v[i]!r}, expected {expected!r}")
        invoices.add(v[1])
    if len(invoices) != rows:
        problems.append("duplicate invoice numbers")
    return problems, rows


def main() -> int:
    print(f"Verifying input files in: {FOLDER}\n")
    failed, total = False, 0
    for name in FILES:
        problems, rows = verify(FOLDER / name)
        total += rows
        if problems:
            failed = True
            print(f"[FAIL] {name}")
            for p in problems[:10]:
                print(f"       - {p}")
        else:
            print(f"[ OK ] {name}: {rows} rows, {len(COLUMNS)} columns, formulas present")
    print(f"\nTotal data rows: {total}")
    print("RESULT: FAILED" if failed else "RESULT: ALL INPUT FILES VERIFIED")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())