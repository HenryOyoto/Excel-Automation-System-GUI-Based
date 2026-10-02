"""Builds Master_Sales_Report.xlsx (five formatted sheets) from consolidated data."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.excel import processor as P
from app.utils import constants as C

CUR, INT, PCT, DATE = "$#,##0.00", "#,##0", "0%", "yyyy-mm-dd"
NAVY, LIGHT = "1F4E78", "DDEBF7"

HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CARD_SIDE = Side(style="medium", color="FFFFFF")
CARD_BORDER = Border(left=CARD_SIDE, right=CARD_SIDE, top=CARD_SIDE, bottom=CARD_SIDE)

CONSOL_FORMATS = {
    C.COL_DATE: DATE, C.COL_QUANTITY: INT, C.COL_UNIT_PRICE: CUR, C.COL_DISCOUNT: PCT,
    C.COL_TAX_RATE: PCT, C.COL_GROSS: CUR, C.COL_DISCOUNT_AMT: CUR,
    C.COL_TAX_AMT: CUR, C.COL_NET: CUR,
}
PRODUCT_FORMATS = {"Transactions": INT, "Total Quantity": INT,
                   "Total Revenue": CUR, "Average Unit Price": CUR}
SALES_FORMATS = {"Transactions": INT, "Quantity Sold": INT,
                 "Total Revenue": CUR, "Average Order Value": CUR}
MONTH_FORMATS = {"Transactions": INT, "Quantity": INT, "Revenue": CUR,
                 "Discount": CUR, "Tax": CUR, "Average Order Value": CUR}
METRIC_FORMATS = {
    "Total Revenue": CUR, "Total Tax": CUR, "Total Discount": CUR, "Average Order Value": CUR,
}  # everything else is an integer count


def _write_table_sheet(wb: Workbook, title: str, df: pd.DataFrame, formats: dict,
                       table_name: str, tab_color: str):
    """Write a DataFrame as a styled Excel Table with frozen header row."""
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab_color
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(row)

    n_rows, n_cols = len(df), len(df.columns)
    for c_idx, name in enumerate(df.columns, start=1):
        head = ws.cell(row=1, column=c_idx)
        head.font, head.fill = HEADER_FONT, HEADER_FILL
        head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        fmt = formats.get(name)
        width = 12 if fmt == DATE else int(df[name].astype(str).str.len().max()) + 3
        width = max(width, len(str(name)) + 3, 12)
        ws.column_dimensions[get_column_letter(c_idx)].width = min(width, 34)

        for r in range(2, n_rows + 2):
            cell = ws.cell(row=r, column=c_idx)
            cell.border = BORDER
            if fmt:
                cell.number_format = fmt

    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    table = Table(displayName=table_name, ref=f"A1:{get_column_letter(n_cols)}{n_rows + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    return ws


def _bar_chart(title, ws_src, data_col, cat_col, last_row, *, horizontal=False,
               width=9.6, height=7.5):
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title = title
    chart.style = 10
    chart.add_data(Reference(ws_src, min_col=data_col, min_row=1, max_row=last_row),
                   titles_from_data=True)
    chart.set_categories(Reference(ws_src, min_col=cat_col, min_row=2, max_row=last_row))
    chart.legend = None
    chart.width, chart.height = width, height
    chart.y_axis.number_format = "$#,##0"
    chart.x_axis.delete = False   # newer Excel hides axes unless told otherwise
    chart.y_axis.delete = False
    if horizontal:
        chart.x_axis.scaling.orientation = "maxMin"  # biggest bar on top
    return chart


def _write_dashboard(ws, df, metrics, ws_prod, ws_sales, ws_month, n_prod, n_sales, n_month):
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 17

    # Title banner
    ws.merge_cells("B2:G2")
    ws["B2"] = "Sales Performance Dashboard"
    ws["B2"].font = Font(size=22, bold=True, color="FFFFFF")
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in "BCDEFG":
        ws[f"{col}2"].fill = HEADER_FILL
    ws.row_dimensions[2].height = 38

    n_files = df[C.COL_SOURCE].nunique()
    span = f"{df[C.COL_DATE].min():%Y-%m-%d} to {df[C.COL_DATE].max():%Y-%m-%d}"
    ws.merge_cells("B3:G3")
    ws["B3"] = (f"Generated {datetime.now():%Y-%m-%d %H:%M}   |   {n_files} source file(s)   |   "
                f"{len(df):,} transactions   |   {span}")
    ws["B3"].font = Font(italic=True, color="595959")

    # KPI cards: 3 per row, each 2 columns wide
    for i, (label, value) in enumerate(metrics.items()):
        r, c = 5 + (i // 3) * 3, 2 + (i % 3) * 2
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 1)
        ws.merge_cells(start_row=r + 1, start_column=c, end_row=r + 1, end_column=c + 1)
        lab, val = ws.cell(row=r, column=c, value=label), ws.cell(row=r + 1, column=c, value=value)
        lab.font = Font(bold=True, size=10, color="44546A")
        val.font = Font(bold=True, size=18, color=NAVY)
        val.number_format = METRIC_FORMATS.get(label, INT)
        for cell in (lab, val):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for rr in (r, r + 1):
            for cc in (c, c + 1):
                ws.cell(row=rr, column=cc).fill = PatternFill("solid", fgColor=LIGHT)
                ws.cell(row=rr, column=cc).border = CARD_BORDER
        ws.row_dimensions[r + 1].height = 36

    # Charts read straight from the analysis sheets
    ws.add_chart(_bar_chart("Revenue by Month", ws_month, 4, 1, n_month + 1), "B14")
    ws.add_chart(_bar_chart("Revenue by Salesperson", ws_sales, 4, 1, n_sales + 1), "E14")
    ws.add_chart(_bar_chart("Top Products by Revenue", ws_prod, 5, 1, min(9, n_prod + 1),
                            horizontal=True, width=19.4, height=8.5), "B30")

    # Traceability: which files feed this report
    ws["B48"] = "Source Files Included"
    ws["B48"].font = Font(bold=True, size=12, color=NAVY)
    heads = {"B": "Source File", "D": "Rows", "E": "Net Sales"}
    for col in "BCDE":
        ws[f"{col}49"].fill, ws[f"{col}49"].font = HEADER_FILL, HEADER_FONT
    for col, text in heads.items():
        ws[f"{col}49"] = text
    per_file = df.groupby(C.COL_SOURCE).agg(rows=(C.COL_INVOICE, "count"), net=(C.COL_NET, "sum"))
    for offset, (name, row) in enumerate(per_file.iterrows(), start=50):
        ws.merge_cells(f"B{offset}:C{offset}")
        ws[f"B{offset}"], ws[f"D{offset}"], ws[f"E{offset}"] = name, int(row["rows"]), float(row["net"])
        ws[f"D{offset}"].number_format, ws[f"E{offset}"].number_format = INT, CUR
        for col in "BCDE":
            ws[f"{col}{offset}"].border = BORDER


def write_master(df: pd.DataFrame, path: Path) -> None:
    """Create the full master workbook at `path` from consolidated data."""
    wb = Workbook()
    ws_dash = wb.active
    ws_dash.title = C.SHEET_DASHBOARD

    products = P.product_analysis(df)
    sales = P.salesperson_analysis(df)
    monthly = P.monthly_summary(df)

    _write_table_sheet(wb, C.SHEET_CONSOLIDATED, df, CONSOL_FORMATS, "ConsolidatedData", "2E75B6")
    ws_prod = _write_table_sheet(wb, C.SHEET_PRODUCTS, products, PRODUCT_FORMATS,
                                 "ProductAnalysis", "548235")
    ws_sales = _write_table_sheet(wb, C.SHEET_SALESPERSONS, sales, SALES_FORMATS,
                                  "SalespersonAnalysis", "BF8F00")
    ws_month = _write_table_sheet(wb, C.SHEET_MONTHLY, monthly, MONTH_FORMATS,
                                  "MonthlySummary", "7030A0")

    # Conditional formatting
    ws_prod.conditional_formatting.add(
        f"E2:E{len(products) + 1}",
        DataBarRule(start_type="min", end_type="max", color="5B9BD5"))
    ws_sales.conditional_formatting.add(
        f"D2:D{len(sales) + 1}",
        ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="63BE7B"))
    ws_month.conditional_formatting.add(
        f"D2:D{len(monthly) + 1}",
        DataBarRule(start_type="min", end_type="max", color="9E7CC1"))

    _write_dashboard(ws_dash, df, P.dashboard_metrics(df), ws_prod, ws_sales, ws_month,
                     len(products), len(sales), len(monthly))
    wb.save(path)